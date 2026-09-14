from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Documentacion_Estado_Actual_VisionClass.docx"
ASSETS = ROOT / "docs" / "_qa_assets"
ASSETS.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEAL = "167D89"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
DARK = "1F2937"
WHITE = "FFFFFF"
GREEN = "176B3A"
GREEN_BG = "EAF5EE"
AMBER = "7A5A00"
AMBER_BG = "FFF4CC"
RED = "9B1C1C"
RED_BG = "FDECEC"

def font(run, name="Calibri", size=11, bold=False, color=DARK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def set_table_geometry(table, widths):
    table.autofit = False
    total = sum(widths)
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths[min(idx, len(widths)-1)]))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        set_cell_shading(hdr.cells[i], LIGHT_BLUE)
        p = hdr.cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(h), size=font_size, bold=True, color=NAVY)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if ridx % 2:
                set_cell_shading(cells[i], "FAFBFC")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            font(p.add_run(str(value)), size=font_size, color=DARK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

def add_callout(doc, title, text, kind="info"):
    colors = {
        "info": (LIGHT_BLUE, NAVY),
        "ok": (GREEN_BG, GREEN),
        "warn": (AMBER_BG, AMBER),
        "risk": (RED_BG, RED),
    }
    fill, color = colors[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(title + "  "), bold=True, color=color)
    font(p.add_run(text), color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    font(p.add_run(text))
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    font(p.add_run(text))
    return p

def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F6F8FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(text.strip().splitlines()):
        r = p.add_run(line)
        font(r, name="Consolas", size=8.7, color="273142")
        if idx < len(text.strip().splitlines()) - 1:
            r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def add_picture_with_alt(run, image_path, width, alt_text):
    picture = run.add_picture(str(image_path), width=width)
    picture._inline.docPr.set("descr", alt_text)
    picture._inline.docPr.set("title", alt_text)
    return picture

def page_break(doc):
    doc.add_page_break()

def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)

def body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead and text.startswith(bold_lead):
        font(p.add_run(bold_lead), bold=True, color=NAVY)
        font(p.add_run(text[len(bold_lead):]))
    else:
        font(p.add_run(text))
    return p

def make_diagram(path, mode):
    img = Image.new("RGB", (1500, 820), "white")
    d = ImageDraw.Draw(img)
    try:
        titlef = ImageFont.truetype("arialbd.ttf", 42)
        boxf = ImageFont.truetype("arialbd.ttf", 27)
        smallf = ImageFont.truetype("arial.ttf", 22)
    except Exception:
        titlef = boxf = smallf = ImageFont.load_default()
    d.text((60, 35), "Arquitectura actual de VisionClass" if mode == "arch" else "Flujo de atención visual", fill="#17365D", font=titlef)
    if mode == "arch":
        boxes = [
            (70, 170, 360, 330, "Usuarios", "Estudiante | Docente | Admin", "#E8EEF5"),
            (470, 120, 810, 300, "Frontend Next.js", "Interfaz, JWT, cámara y proxy ML", "#DDEBF7"),
            (470, 410, 810, 590, "Backend Django/DRF", "API, permisos, métricas y reportes", "#E2F0D9"),
            (930, 120, 1360, 300, "Servicio ML FastAPI", "MediaPipe/OpenCV; ONNX opcional", "#FFF2CC"),
            (930, 410, 1360, 590, "PostgreSQL", "Cursos, sesiones, eventos y resultados", "#FCE4D6"),
        ]
        for x1,y1,x2,y2,t,s,c in boxes:
            d.rounded_rectangle((x1,y1,x2,y2), radius=18, fill=c, outline="#4B5563", width=3)
            d.text((x1+20,y1+30),t,fill="#17365D",font=boxf)
            d.multiline_text((x1+20,y1+85),s,fill="#374151",font=smallf,spacing=7)
        arrows=[((360,250),(470,210)),((640,300),(640,410)),((810,210),(930,210)),((810,500),(930,500)),((1145,300),(1145,410))]
    else:
        boxes = [
            (55, 260, 290, 450, "1. Cámara", "Frame JPEG cada ~1 s", "#E8EEF5"),
            (350, 260, 610, 450, "2. Proxy", "Valida JWT y rol", "#DDEBF7"),
            (670, 260, 960, 450, "3. Visión", "Rostro, ojos y mirada", "#FFF2CC"),
            (1020, 260, 1320, 450, "4. Persistencia", "Evento y promedio", "#E2F0D9"),
        ]
        for x1,y1,x2,y2,t,s,c in boxes:
            d.rounded_rectangle((x1,y1,x2,y2), radius=18, fill=c, outline="#4B5563", width=3)
            d.text((x1+18,y1+35),t,fill="#17365D",font=boxf)
            d.multiline_text((x1+18,y1+95),s,fill="#374151",font=smallf,spacing=7)
        arrows=[((290,355),(350,355)),((610,355),(670,355)),((960,355),(1020,355))]
        d.rounded_rectangle((675,560,1315,710), radius=16, fill="#FDECEC", outline="#9B1C1C", width=3)
        d.text((700,585),"Estado real",fill="#9B1C1C",font=boxf)
        d.multiline_text((700,635),"Heurística operativa; CNN-LSTM preparado pero sin artefacto ONNX versionado.",fill="#374151",font=smallf,spacing=6)
    for (a,b) in arrows:
        d.line((a,b), fill="#167D89", width=7)
        x,y=b
        d.polygon([(x,y),(x-16,y-10),(x-16,y+10)],fill="#167D89")
    img.save(path, quality=95)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# Styles: standard_business_brief with editorial_cover opening.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_name in ["List Bullet", "List Bullet 2", "List Number"]:
    st = styles[list_name]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(4)

# Header/footer
hp = section.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
font(hp.add_run("VISIONCLASS  |  DOCUMENTACIÓN DEL ESTADO ACTUAL"), size=8.5, bold=True, color=MID_GRAY)
fp = section.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(fp.add_run("Documento técnico  |  "), size=8.5, color=MID_GRAY)
fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
r = fp.add_run(); r._r.append(fld_begin); r._r.append(instr); r._r.append(fld_end)

# Cover
for _ in range(4): doc.add_paragraph()
logo = ROOT / "frontend" / "public" / "LOGO.png"
if logo.exists():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture_with_alt(p.add_run(), logo, Inches(1.55), "Logotipo de VisionClass")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("INFORME TÉCNICO"), size=11, bold=True, color=TEAL)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
font(p.add_run("Documentación del estado actual"), size=29, bold=True, color=NAVY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(26)
font(p.add_run("VisionClass - Plataforma educativa con analítica de atención visual"), size=14, color=DARK_BLUE)
add_table(doc, ["Corte de revisión", "Versión documental", "Repositorio auditado"], [["14 de agosto de 2026", "1.0", "Commit 7f8f9cb"]], [3000, 2600, 3760], 9.5)
for _ in range(4): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("Documento de referencia para desarrollo, operación, evaluación y escalamiento"), size=10, italic=True, color=MID_GRAY)
page_break(doc)

add_heading(doc, "Control del documento", 1)
add_table(doc, ["Campo", "Valor"], [
    ["Título", "Documentación del estado actual de VisionClass"],
    ["Propósito", "Consolidar el estado técnico y funcional verificable de la aplicación."],
    ["Audiencia", "Equipo de desarrollo, responsables académicos, operación, seguridad y futuros mantenedores."],
    ["Método", "Revisión estática del repositorio, configuración, rutas, modelos, despliegue y flujo ML."],
    ["Limitación", "No incluye mediciones directas de disponibilidad, latencia ni datos internos del entorno Render."],
], [2300, 7060])
add_heading(doc, "Historial de versiones", 2)
add_table(doc, ["Versión", "Fecha", "Cambio", "Responsable"], [["1.0", "14/08/2026", "Emisión inicial basada en el commit 7f8f9cb.", "Equipo VisionClass"]], [1200, 1600, 4560, 2000])
add_callout(doc, "Clasificación del contenido", "Las credenciales, claves y secretos se representan únicamente como nombres de variables. Este documento no reproduce valores sensibles de archivos .env.", "info")

add_heading(doc, "Contenido", 1)
toc_items = [
    "1. Resumen ejecutivo", "2. Alcance y metodología", "3. Visión funcional",
    "4. Arquitectura de solución", "5. Componentes técnicos", "6. Datos y API",
    "7. Flujos principales", "8. Estado del machine learning", "9. Seguridad y privacidad",
    "10. Despliegue y operación", "11. Calidad y mantenibilidad", "12. Riesgos y deuda técnica",
    "13. Escalabilidad", "14. Hoja de ruta", "15. Runbook operativo", "Apéndices"
]
for item in toc_items: add_bullet(doc, item)
page_break(doc)

add_heading(doc, "1. Resumen ejecutivo", 1)
body(doc, "VisionClass es una aplicación web educativa compuesta por un frontend Next.js, una API Django REST Framework, una base PostgreSQL y un servicio independiente FastAPI para análisis visual. La aplicación ofrece autenticación local y Google OAuth, gestión de cursos y contenidos, inscripción, consumo de materiales, cuestionarios finales, prueba D2R, seguimiento de atención visual, paneles por rol, notificaciones y administración institucional.")
body(doc, "El estado general se clasifica como prototipo funcional desplegable. Los recorridos principales están representados en código y existe configuración para Render y Docker Compose. No obstante, la madurez es desigual: las capacidades LMS y de administración están más consolidadas que la validación científica del indicador de atención, la automatización de pruebas y la preparación para concurrencia elevada.")
add_table(doc, ["Área", "Estado", "Evaluación resumida"], [
    ["Frontend", "Operativo", "Rutas y paneles por rol; diseño responsive incorporado en el último commit."],
    ["Backend/API", "Operativo", "Recursos REST, JWT, OAuth, reportes, permisos y persistencia relacional."],
    ["Servicio ML", "Parcial", "Visión computacional heurística operativa; CNN-LSTM opcional sin artefacto versionado."],
    ["D2R", "Operativo con cautelas", "Captura resultados y atención; no debe presentarse como diagnóstico clínico."],
    ["Despliegue", "Operativo", "Blueprint Render y entorno Docker local disponibles."],
    ["Pruebas", "Insuficiente", "Cobertura automatizada prácticamente ausente en backend y ML."],
    ["Escalabilidad", "En riesgo", "El envío por estudiante de ~1 frame/s satura CPU y puede acumular solicitudes."],
], [1900, 1800, 5660])
add_callout(doc, "Conclusión ejecutiva", "La aplicación demuestra valor y una arquitectura modular, pero antes de escalar debe cerrar cuatro brechas: validación ML, control de concurrencia, observabilidad y pruebas automatizadas.", "warn")

add_heading(doc, "2. Alcance y metodología", 1)
body(doc, "La revisión cubre el código fuente presente en frontend, backend y ml; los archivos de despliegue render.yaml y docker-compose.yml; las dependencias; las rutas públicas; los modelos persistentes; los comandos de gestión y la documentación existente. El corte corresponde al commit 7f8f9cb, cuyo mensaje indica mejoras responsive y de contenido.")
add_heading(doc, "2.1 Incluido", 2)
for x in ["Inventario funcional por actor.", "Arquitectura lógica y flujo de datos.", "Estado real del módulo de atención visual y entrenamiento.", "Modelo de datos y superficie API.", "Configuración, despliegue, seguridad, operación y mantenimiento.", "Riesgos, prioridades y hoja de ruta."]: add_bullet(doc, x)
add_heading(doc, "2.2 Fuera de alcance", 2)
for x in ["Auditoría de infraestructura activa en Render.", "Prueba de penetración y revisión legal formal.", "Validación psicométrica o clínica del D2R digital.", "Métricas de producción no almacenadas en el repositorio.", "Certificación de accesibilidad WCAG completa."]: add_bullet(doc, x)

add_heading(doc, "3. Visión funcional", 1)
add_heading(doc, "3.1 Actores y capacidades", 2)
add_table(doc, ["Actor", "Capacidades actuales"], [
    ["Estudiante", "Autenticarse, explorar e inscribirse, consumir materiales, habilitar cámara, realizar quiz final y D2R, consultar métricas y notificaciones."],
    ["Docente", "Crear y organizar cursos, módulos, lecciones y materiales; consultar estadísticas; revisar resultados D2R; enviar notificaciones."],
    ["Administrador", "Gestionar usuarios y cursos, consultar analítica agregada, administrar políticas de privacidad y solicitudes de investigación."],
    ["Servicios externos", "Google OAuth para identidad, Mailgun para correo y Google GenAI para generación asistida de pruebas."],
], [1800, 7560])
add_heading(doc, "3.2 Mapa de funcionalidades", 2)
add_table(doc, ["Dominio", "Funcionalidad", "Estado"], [
    ["Identidad", "JWT por email/usuario y Google OAuth", "Operativo"],
    ["Cursos", "Curso > módulo > lección > material", "Operativo"],
    ["Inscripción", "Inscripción, estado y progreso en JSON", "Operativo"],
    ["Evaluación", "Quiz final y registro de intentos", "Operativo"],
    ["Atención", "Captura, análisis por frame y agregación", "Parcial"],
    ["D2R", "Sesiones, eventos, resultados y programación", "Operativo"],
    ["Reportes", "Paneles y exportación PDF/CSV/XLSX", "Operativo"],
    ["Notificaciones", "Persistencia y envío por Mailgun", "Condicionado a configuración"],
    ["Administración", "Usuarios, cursos, analítica, privacidad e investigación", "Operativo"],
], [1600, 5600, 2160])
page_break(doc)

add_heading(doc, "4. Arquitectura de solución", 1)
arch = ASSETS / "architecture.png"; make_diagram(arch, "arch")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_picture_with_alt(
    p.add_run(),
    arch,
    Inches(6.3),
    "Diagrama de arquitectura lógica actual de VisionClass",
)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("Figura 1. Arquitectura lógica verificada en el repositorio."), size=9, italic=True, color=MID_GRAY)
add_heading(doc, "4.1 Responsabilidades", 2)
add_table(doc, ["Componente", "Tecnología", "Responsabilidad"], [
    ["Frontend", "Next.js 16, React 19, TypeScript", "Experiencia de usuario, cámara, estado de sesión, paneles y proxy ML."],
    ["Backend", "Django 5.2, DRF 3.16", "Reglas de negocio, autenticación, autorización, API, reportes y administración."],
    ["ML", "FastAPI, MediaPipe, OpenCV, ONNX Runtime", "Detección facial, heurística de atención, ventana temporal e inferencia opcional."],
    ["Datos", "PostgreSQL 15", "Persistencia de usuarios, cursos, sesiones, eventos, resultados y políticas."],
    ["Integraciones", "Google, Mailgun, Google GenAI", "Autenticación social, correo y generación de cuestionarios."],
], [1600, 2600, 5160])
add_heading(doc, "4.2 Características arquitectónicas", 2)
for x in ["Separación clara entre interfaz, negocio e inferencia.", "Comunicación HTTP/JSON y multipart entre servicios.", "Persistencia centralizada en PostgreSQL.", "Frontend actúa como proxy autenticado hacia ML.", "Despliegue independiente de cada servicio en Render."]: add_bullet(doc, x)
add_callout(doc, "Observación", "El estado temporal del ML se conserva en memoria del proceso. Esto impide una escalabilidad horizontal segura sin afinidad de sesión o almacenamiento compartido.", "risk")

add_heading(doc, "5. Componentes técnicos", 1)
add_heading(doc, "5.1 Frontend", 2)
body(doc, "El frontend utiliza App Router y dispone de páginas para landing, login, estudiante, curso, instructor, administrador, D2R y contenidos legales. Emplea Tailwind CSS, Radix UI, Recharts, Lucide y componentes reutilizables. La autenticación se conserva en contexto y las llamadas se concentran en app/lib/api.ts.")
for x in ["Rutas protegidas diferenciadas por rol.", "Captura de cámara con getUserMedia y canvas.", "Proxy servidor /api/attention-proxy para no exponer directamente toda la lógica del servicio ML.", "Paneles con métricas, gráficos, filtros y exportaciones.", "Componentes responsive ajustados en el commit auditado."]: add_bullet(doc, x)
add_heading(doc, "5.2 Backend", 2)
body(doc, "La API usa un usuario personalizado con roles student, teacher y admin. DRF aplica autenticación JWT de forma predeterminada. Los ViewSets filtran recursos según el usuario y los endpoints administrativos requieren rol o privilegios elevados. WhiteNoise sirve estáticos y dj-database-url habilita PostgreSQL administrado.")
add_heading(doc, "5.3 Servicio ML", 2)
body(doc, "El servicio carga MediaPipe Face Mesh cuando está disponible y conserva Haar Cascade como respaldo. Para cada imagen calcula presencia facial, apertura ocular y desviación de mirada; luego agrega una ventana temporal. El modelo CNN-LSTM se intenta cargar desde MODEL_PATH, pero el repositorio no contiene un archivo ONNX o métricas versionadas.")
add_heading(doc, "5.4 Base de datos", 2)
body(doc, "PostgreSQL es la fuente de verdad. Las relaciones estructuran cursos, actividad y analítica. Parte del progreso y metadatos se guarda en JSONField, lo que facilita evolución rápida pero exige contratos y validación para evitar esquemas implícitos divergentes.")

add_heading(doc, "6. Datos y API", 1)
add_heading(doc, "6.1 Modelo de datos", 2)
add_table(doc, ["Entidad", "Propósito", "Relaciones principales"], [
    ["User", "Identidad, rol y perfil", "Cursos, inscripciones, sesiones y resultados"],
    ["Course", "Unidad educativa principal", "Owner, módulos, inscripciones y sesiones"],
    ["CourseModule", "Agrupación ordenada", "Curso y lecciones"],
    ["CourseLesson", "Lección ordenada", "Módulo y materiales"],
    ["CourseMaterial", "PDF, video, enlace, quiz u otro", "Lección y metadata"],
    ["Enrollment", "Vínculo usuario-curso", "Estado y enrollment_data"],
    ["Session", "Actividad de curso", "Usuario, curso, eventos y vistas"],
    ["AttentionEvent", "Muestra de atención", "Sesión, usuario, tiempo y datos"],
    ["D2RSession / Result", "Ejecución y resultado D2R", "Usuario, eventos y phase_data"],
    ["ContentView", "Consumo de contenido", "Sesión, usuario, duración"],
    ["QuizAttempt", "Intento de evaluación", "Sesión, usuario, score"],
    ["Notification", "Mensaje y estado de entrega", "Remitente, destinatario y curso"],
    ["Privacy / Research", "Gobierno de datos", "Políticas y solicitudes"],
], [1900, 3700, 3760], 8.8)
add_heading(doc, "6.2 Superficie API principal", 2)
add_table(doc, ["Ruta", "Uso", "Acceso"], [
    ["/api/auth/token/", "Inicio de sesión local", "Público"],
    ["/api/auth/token/refresh/", "Renovación JWT", "Refresh token"],
    ["/api/auth/google/", "Intercambio OAuth", "Público controlado"],
    ["/api/me/", "Perfil y rol actual", "Autenticado"],
    ["/api/courses/", "Gestión y consulta de cursos", "Según rol"],
    ["/api/course-modules/", "Módulos", "Según pertenencia"],
    ["/api/course-lessons/", "Lecciones", "Según pertenencia"],
    ["/api/course-materials/", "Materiales y descarga", "Según pertenencia"],
    ["/api/enrollments/", "Inscripciones y progreso", "Autenticado"],
    ["/api/sessions/", "Sesiones de curso", "Autenticado"],
    ["/api/attention-events/", "Eventos de atención", "Autenticado/servicio"],
    ["/api/d2r-results/", "Resultados D2R", "Según rol"],
    ["/api/student-metrics/", "Panel del estudiante", "Estudiante"],
    ["/api/exports/student-report/", "PDF, CSV o XLSX", "Autenticado"],
    ["/api/ai/generate-test/", "Generación asistida", "Docente"],
    ["/api/admin/overview/", "Resumen institucional", "Admin"],
    ["/api/admin/analytics/", "Analítica agregada", "Admin"],
], [3000, 4160, 2200], 8.7)
page_break(doc)

add_heading(doc, "7. Flujos principales", 1)
add_heading(doc, "7.1 Autenticación", 2)
for x in ["El usuario ingresa credenciales o inicia Google OAuth.", "El backend valida identidad y emite access/refresh JWT.", "El frontend consulta /api/me/ para conocer rol y redirigir.", "Las solicitudes posteriores incluyen Authorization: Bearer."]: add_number(doc, x)
add_heading(doc, "7.2 Inscripción y consumo", 2)
for x in ["El estudiante selecciona un curso y crea Enrollment.", "Se presenta la estructura curso-módulo-lección-material.", "Al consumir contenido se crea Session y ContentView.", "El progreso y atención agregada actualizan enrollment_data.", "El quiz final puede marcar el curso como completado."]: add_number(doc, x)
add_heading(doc, "7.3 Atención visual", 2)
flow = ASSETS / "ml_flow.png"; make_diagram(flow, "ml")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_picture_with_alt(
    p.add_run(),
    flow,
    Inches(6.3),
    "Flujo actual de análisis de atención visual de VisionClass",
)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("Figura 2. Flujo actual de análisis de atención."), size=9, italic=True, color=MID_GRAY)
add_heading(doc, "7.4 D2R", 2)
body(doc, "La prueba D2R es un módulo independiente de los cursos. Registra sesión, respuestas por fase, velocidad, aciertos, errores y eventos de atención. El docente puede consultar el último resultado. Los indicadores derivados deben presentarse como apoyo educativo y no como diagnóstico clínico.")

add_heading(doc, "8. Estado del machine learning", 1)
add_callout(doc, "Evaluación honesta", "Existe procesamiento real de visión computacional. Sin embargo, la inferencia productiva no puede atribuirse hoy a un modelo entrenado mientras el artefacto ONNX y sus métricas no formen parte del despliegue verificable.", "warn")
add_heading(doc, "8.1 Lo que sí está implementado", 2)
for x in ["Decodificación de imágenes con OpenCV.", "Face Mesh e iris de MediaPipe cuando inicializa correctamente.", "Fallback Haar para rostro y ojos.", "Puntuación heurística entre 0 y 1.", "Ventana temporal por sesión.", "Persistencia de eventos en backend.", "Pipeline de entrenamiento CNN-LSTM y exportación ONNX."]: add_bullet(doc, x)
add_heading(doc, "8.2 Lo experimental", 2)
for x in ["No existe cnn_lstm.onnx versionado en el repositorio.", "No existe metrics.json con resultados reproducibles.", "La función aggregate_temporal_score está descrita como placeholder.", "No hay evidencia de calibración con observadores humanos o población normativa.", "XGBoost, scikit-learn y joblib están instalados, pero no forman parte clara de la inferencia activa."]: add_bullet(doc, x)
add_heading(doc, "8.3 Arquitectura de entrenamiento preparada", 2)
add_table(doc, ["Etapa", "Implementación actual", "Brecha"], [
    ["Datos", "Exportación de secuencias D2R a Parquet", "Requiere frames guardados y gobierno de consentimiento"],
    ["Separación", "Split por usuario", "Documentar semilla y composición"],
    ["Modelo", "MobileNetV3 Small + LSTM + Sigmoid", "Validar tamaño, latencia y sobreajuste"],
    ["Optimización", "Huber Loss y Adam", "Registrar hiperparámetros y experimentos"],
    ["Evaluación", "Loss y MAE", "Añadir calibración, correlación, intervalos y sesgo"],
    ["Entrega", "Exportación ONNX", "Versionado, checksum, registry y rollback"],
], [1600, 3700, 4060])
add_heading(doc, "8.4 Lenguaje recomendado", 2)
body(doc, "En presentaciones y producto debe utilizarse “estimación de atención visual” o “indicador de interacción visual”. Debe evitarse afirmar que el sistema mide el estado cognitivo, diagnostica atención o determina aprendizaje por sí solo.")

add_heading(doc, "9. Seguridad y privacidad", 1)
add_heading(doc, "9.1 Controles existentes", 2)
for x in ["JWT con expiración configurable.", "OAuth de Google mediante django-allauth.", "CORS y CSRF restringibles por variable.", "Redirección HTTPS y cookies seguras en Render.", "Permisos por rol y endpoints administrativos dedicados.", "SAVE_FRAMES desactivado por defecto.", "Políticas y solicitudes de investigación persistidas."]: add_bullet(doc, x)
add_heading(doc, "9.2 Riesgos", 2)
add_table(doc, ["Riesgo", "Severidad", "Tratamiento recomendado"], [
    ["Contraseñas locales de ejemplo en Docker", "Alta", "Usar secretos y eliminar valores conocidos."],
    ["Servicio ML acepta identificadores del formulario", "Alta", "Vincular user_id al JWT validado y firmar tráfico interno."],
    ["Fallback SECRET_KEY inseguro", "Alta", "Fallar el arranque en producción si falta la variable."],
    ["Datos biométricos/visuales", "Alta", "Consentimiento, minimización, retención y DPIA."],
    ["Logs detallados", "Media", "Evitar payloads, tokens e identificadores innecesarios."],
    ["Dependencias sin política de actualización", "Media", "SCA, lockfiles y calendario de parches."],
], [3200, 1300, 4860], 8.8)
add_callout(doc, "Privacidad por diseño", "Mantener SAVE_FRAMES=0 en producción salvo un protocolo aprobado. Si se capturan imágenes para entrenamiento, separar propósito académico, consentimiento, cifrado, acceso y plazo de eliminación.", "risk")

add_heading(doc, "10. Despliegue y operación", 1)
add_heading(doc, "10.1 Render", 2)
body(doc, "render.yaml define PostgreSQL administrado y tres servicios web independientes. Backend ejecuta migraciones y collectstatic antes del despliegue; frontend compila Next.js; ML inicia Uvicorn. Las URL públicas están configuradas para los nombres visionclass-backend, visionclass-frontend y visionclass-ml.")
add_table(doc, ["Servicio", "Build", "Start", "Dependencias clave"], [
    ["Backend", "pip install", "gunicorn core.wsgi", "PostgreSQL, OAuth, Mailgun, GenAI"],
    ["Frontend", "npm install && npm run build", "npm run start", "Backend y ML"],
    ["ML", "pip install", "uvicorn ml_service:app", "Backend, MediaPipe, OpenCV, ONNX"],
], [1500, 2500, 2500, 2860])
add_heading(doc, "10.2 Docker local", 2)
add_code(doc, "docker compose up -d --build\ndocker compose exec backend python manage.py migrate\ndocker compose exec backend python manage.py create_default_admin")
add_heading(doc, "10.3 Variables esenciales", 2)
add_table(doc, ["Componente", "Variables"], [
    ["Backend", "SECRET_KEY, DATABASE_URL, ALLOWED_HOSTS, CORS_ALLOWED_ORIGINS, CSRF_TRUSTED_ORIGINS, GOOGLE_OAUTH_*, MAILGUN_*, GOOGLE_API_KEY"],
    ["Frontend", "NEXT_PUBLIC_BACKEND_URL, NEXT_PUBLIC_ML_URL, ML_SERVICE_URL, NEXT_PUBLIC_GOOGLE_CLIENT_ID, NEXT_PUBLIC_GOOGLE_REDIRECT_URI"],
    ["ML", "BACKEND_URL, BACKEND_TOKEN, SEQUENCE_LENGTH, MODEL_PATH, MODEL_IMG_SIZE, CORS_ORIGINS, SAVE_FRAMES, FRAMES_DIR"],
], [1800, 7560], 9.0)
add_heading(doc, "10.4 Health checks", 2)
add_code(doc, "curl -i https://visionclass-backend.onrender.com/api/me/\ncurl -i https://visionclass-ml.onrender.com/health\ncurl -I https://visionclass-frontend.onrender.com/")
body(doc, "El primer endpoint requiere JWT y por ello puede devolver 401, lo cual confirma que la aplicación responde. Para operación formal conviene crear /api/health/ público y mínimo, sin datos sensibles, con comprobación opcional de base de datos.")

add_heading(doc, "11. Calidad y mantenibilidad", 1)
add_table(doc, ["Dimensión", "Estado actual", "Objetivo"], [
    ["Pruebas unitarias", "Archivo backend casi vacío; sin suite ML", "Cobertura de negocio y visión"],
    ["Pruebas integración", "No automatizadas", "Flujos OAuth, inscripción, curso y D2R"],
    ["CI/CD", "Cloud Build y Render presentes", "Gate de lint, test, build y migración"],
    ["Observabilidad", "Logs con print/console", "Logs estructurados, métricas y alertas"],
    ["Documentación", "README y notas de diagnóstico", "Documento vivo y ADRs"],
    ["Dependencias", "Versiones mixtas fijas y abiertas", "Pinning y actualización controlada"],
], [1900, 3600, 3860])
add_heading(doc, "11.1 Pruebas mínimas prioritarias", 2)
for x in ["Permisos y aislamiento de datos por rol.", "Inscripción idempotente y finalización de curso.", "CRUD jerárquico de cursos.", "OAuth, refresh y expiración JWT.", "ML con rostro, sin rostro, imagen corrupta y timeout.", "Carga concurrente con 15, 30 y 50 estudiantes.", "Exportaciones PDF/CSV/XLSX.", "Migrations desde base limpia y versión anterior."]: add_bullet(doc, x)

add_heading(doc, "12. Riesgos y deuda técnica", 1)
add_table(doc, ["ID", "Hallazgo", "Impacto", "Prioridad"], [
    ["R-01", "Modelo ONNX y métricas no versionados", "No se demuestra inferencia ML entrenada", "Crítica"],
    ["R-02", "1 frame/s por estudiante con setInterval", "Acumulación y saturación de CPU", "Crítica"],
    ["R-03", "Buffers ML en memoria sin limpieza", "Pérdida, fuga y mala escala horizontal", "Alta"],
    ["R-04", "Sin suite de pruebas", "Regresiones y despliegues inseguros", "Alta"],
    ["R-05", "Indicadores D2R no baremados", "Interpretaciones académicas incorrectas", "Alta"],
    ["R-06", "JSONField para contratos críticos", "Datos inconsistentes", "Media"],
    ["R-07", "Dependencias ML pesadas no utilizadas", "Build lento y superficie mayor", "Media"],
    ["R-08", "Ausencia de observabilidad", "Diagnóstico tardío", "Alta"],
], [900, 3800, 3160, 1500], 8.6)

add_heading(doc, "13. Escalabilidad", 1)
body(doc, "Con la frecuencia actual, 15 estudiantes pueden generar aproximadamente 15 solicitudes de imagen por segundo. Si cada análisis tarda más de un segundo, setInterval puede iniciar otra solicitud antes de finalizar la anterior. La prueba previa con 15 alumnos y una instancia de 0.5 CPU es coherente con saturación y timeouts.")
add_heading(doc, "13.1 Medidas inmediatas", 2)
for x in ["Implementar un candado client-side: no enviar si existe una inferencia pendiente.", "Reducir frecuencia a 2-5 segundos o adaptarla a latencia.", "Reducir resolución JPEG y recortar rostro en cliente cuando sea viable.", "Configurar al menos 1 CPU y 2 GB RAM para pruebas; medir antes de asumir capacidad.", "Añadir límites, timeout, reintentos con jitter y circuit breaker.", "Limpiar buffers al cerrar sesión."]: add_bullet(doc, x)
add_heading(doc, "13.2 Evolución de arquitectura", 2)
for x in ["Cola asíncrona para frames y backpressure.", "Redis para estado temporal y afinidad de sesión.", "Réplicas ML stateless con autoscaling.", "Modelo ligero optimizado y cuantizado.", "Métricas p50/p95/p99, FPS efectivo, no-face ratio y CPU/RAM.", "Separación de ingesta, inferencia y persistencia."]: add_bullet(doc, x)

add_heading(doc, "14. Hoja de ruta recomendada", 1)
add_table(doc, ["Horizonte", "Objetivos", "Criterio de salida"], [
    ["0-30 días", "Pruebas críticas, health checks, candado de envío, logs estructurados, secretos", "Build reproducible y prueba de 15 usuarios sin timeout"],
    ["31-60 días", "Entrenar/versionar ONNX, métricas, model card, monitoreo", "Modelo cargado verificablemente y evaluación por usuario"],
    ["61-90 días", "Redis/cola, load testing, optimización, políticas de datos", "30-50 usuarios con SLO definido"],
    ["3-6 meses", "Validación externa, sesgo, accesibilidad, seguridad formal", "Piloto institucional gobernado"],
], [1500, 5000, 2860], 9.0)
add_heading(doc, "14.1 Definition of Done para ML", 2)
for x in ["Artefacto ONNX con versión y checksum.", "Dataset y consentimiento documentados.", "Separación por usuario y semilla reproducible.", "Métricas de baseline y modelo en conjunto no visto.", "Pruebas de latencia y memoria en CPU.", "Model card con limitaciones y grupos evaluados.", "Rollback y fallback explícito observables."]: add_bullet(doc, x)

add_heading(doc, "15. Runbook operativo", 1)
add_heading(doc, "15.1 Incidente: ML no responde", 2)
for x in ["Confirmar /health y /debug/status.", "Revisar CPU, RAM, reinicios y latencia en Render.", "Verificar mediapipe_initialized y onnx_model_loaded.", "Reducir temporalmente frecuencia de envío.", "Reiniciar servicio si existe fuga de buffers.", "Registrar duración, usuarios afectados y causa raíz."]: add_number(doc, x)
add_heading(doc, "15.2 Incidente: backend 500", 2)
for x in ["Correlacionar hora, ruta, usuario y despliegue.", "Consultar logs de Gunicorn/Django.", "Verificar migraciones y conexión PostgreSQL.", "Reproducir con payload anonimizado.", "Aplicar corrección, prueba y rollback si procede."]: add_number(doc, x)
add_heading(doc, "15.3 Mantenimiento", 2)
add_table(doc, ["Frecuencia", "Actividad"], [
    ["Diaria", "Errores 5xx, disponibilidad, CPU/RAM y cola de notificaciones."],
    ["Semanal", "Backups, crecimiento de eventos, dependencias y usuarios anómalos."],
    ["Mensual", "Restauración de backup, rotación de secretos y revisión de costes."],
    ["Por release", "Tests, migraciones, model version, changelog y smoke test."],
], [1800, 7560])

page_break(doc)
add_heading(doc, "Apéndice A. Inventario tecnológico", 1)
add_table(doc, ["Capa", "Tecnologías principales"], [
    ["Frontend", "Next.js 16.0.1; React 19.2; TypeScript 5; Tailwind 4; Radix UI; Recharts"],
    ["Backend", "Python; Django 5.2.7; DRF 3.16.1; SimpleJWT; allauth; Gunicorn"],
    ["ML", "FastAPI; Uvicorn; OpenCV 4.8; MediaPipe 0.10; PyTorch 2.1; ONNX Runtime"],
    ["Datos", "PostgreSQL 15; JSONField; migraciones Django"],
    ["Operación", "Docker Compose; Render Blueprint; Cloud Build"],
    ["Externos", "Google OAuth; Google GenAI; Mailgun"],
], [1800, 7560])

add_heading(doc, "Apéndice B. Comandos de administración", 1)
add_code(doc, "# Crear administrador si no existe\npython manage.py create_default_admin\n\n# Crear usuario técnico ML\npython manage.py create_ml_service_user\n\n# Migraciones y estáticos\npython manage.py migrate\npython manage.py collectstatic --noinput\n\n# Exportar dataset D2R\npython manage.py export_d2r_frames_dataset --out data/frames_dataset.parquet")
add_callout(doc, "Seguridad", "Definir ADMIN_USERNAME, ADMIN_EMAIL y ADMIN_PASSWORD mediante secretos. No conservar la contraseña predeterminada en producción.", "risk")

add_heading(doc, "Apéndice C. Criterios de aceptación para el próximo release", 1)
criteria = [
    "El frontend compila sin errores TypeScript y supera lint.",
    "El backend inicia desde una base vacía y aplica todas las migraciones.",
    "Los tres roles solo acceden a los recursos autorizados.",
    "La inscripción es idempotente y muestra confirmación clara.",
    "El quiz final completa el curso según la regla aprobada.",
    "El servicio ML no acepta más de una solicitud simultánea por sesión.",
    "El estado /debug/status identifica método y versión de modelo.",
    "Las pruebas de 15 usuarios no generan timeouts ni reinicios.",
    "Los frames no se almacenan sin consentimiento explícito.",
    "Existe evidencia de backup y restauración de PostgreSQL.",
]
for x in criteria: add_bullet(doc, "[ ] " + x)

add_heading(doc, "Apéndice D. Fuentes internas revisadas", 1)
for x in [
    "README.md y documentación de diagnóstico.",
    "render.yaml y docker-compose.yml.",
    "backend/core/settings.py, backend/api/models.py, serializers.py, views.py y urls.py.",
    "frontend/package.json, rutas App Router, contexto de autenticación y proxy de atención.",
    "ml/ml_service.py, train_model.py, train_cnn_lstm.py y requirements.txt.",
    "Migraciones y comandos de gestión Django.",
]: add_bullet(doc, x)
body(doc, "Este informe describe el estado verificable del código en el corte indicado. Debe actualizarse después de cambios de arquitectura, contratos API, política de datos, modelo ML o infraestructura.")

# Document properties
doc.core_properties.title = "Documentación del estado actual de VisionClass"
doc.core_properties.subject = "Arquitectura, funcionalidades, ML, seguridad, despliegue y hoja de ruta"
doc.core_properties.author = "Equipo VisionClass"
doc.core_properties.keywords = "VisionClass, LMS, Django, Next.js, FastAPI, Machine Learning, D2R"
doc.core_properties.comments = "Generado a partir de una revisión estática del commit 7f8f9cb."

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
